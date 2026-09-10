"""Safe, bounded extraction of files into untrusted security evidence."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Optional

from aegis.app.utils.redaction import redact_secrets


@dataclass(frozen=True)
class ProcessedEvidence:
    source_name: str
    media_type: str
    evidence_type: str
    content_hash: str
    text: str
    metadata: dict[str, Any]


class EvidenceProcessor:
    """Convert supported evidence files without executing their content."""

    _TEXT_EXTENSIONS = {".txt", ".md", ".log", ".yaml", ".yml"}
    _JSON_EXTENSIONS = {".json", ".jsonl"}
    _CSV_EXTENSIONS = {".csv", ".tsv"}
    _DOCUMENT_EXTENSIONS = {".pdf", ".docx"}
    _IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp"}
    _MAX_DOCX_ENTRIES = 5000
    _MAX_DOCX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
    _MAX_DOCX_COMPRESSION_RATIO = 200

    def __init__(
        self,
        *,
        max_bytes: int = 20 * 1024 * 1024,
        max_text_chars: int = 200_000,
        image_extractor: Optional[Callable[[bytes], str]] = None,
    ) -> None:
        self.max_bytes = max(1024, int(max_bytes))
        self.max_text_chars = max(1000, int(max_text_chars))
        self._image_extractor = image_extractor

    def process(self, filename: str, data: bytes, media_type: str = "") -> ProcessedEvidence:
        source_name = Path(filename).name[:255]
        suffix = Path(source_name).suffix.lower()
        if not data:
            raise ValueError("evidence file is empty")
        if len(data) > self.max_bytes:
            raise ValueError("evidence file exceeds size limit")
        allowed = (
            self._TEXT_EXTENSIONS
            | self._JSON_EXTENSIONS
            | self._CSV_EXTENSIONS
            | self._DOCUMENT_EXTENSIONS
            | self._IMAGE_EXTENSIONS
        )
        if suffix not in allowed:
            raise ValueError(f"unsupported evidence type: {suffix or media_type}")

        extraction = "text"
        if suffix in self._TEXT_EXTENSIONS:
            text = self._decode_text(data)
        elif suffix in self._JSON_EXTENSIONS:
            text = self._extract_json(data, json_lines=suffix == ".jsonl")
            extraction = "structured_json"
        elif suffix in self._CSV_EXTENSIONS:
            text = self._extract_csv(data, delimiter="\t" if suffix == ".tsv" else ",")
            extraction = "structured_table"
        elif suffix == ".pdf":
            text = self._extract_pdf(data)
            extraction = "pdf_text"
        elif suffix == ".docx":
            text = self._extract_docx(data)
            extraction = "docx_text"
        else:
            text = self._extract_image(data)
            extraction = "vision_or_ocr"

        redacted, redaction_events = redact_secrets(text[: self.max_text_chars], direction="incoming")
        evidence_type = self._evidence_type(suffix)
        return ProcessedEvidence(
            source_name=source_name,
            media_type=media_type or self._default_media_type(suffix),
            evidence_type=evidence_type,
            content_hash=hashlib.sha256(data).hexdigest(),
            text=str(redacted),
            metadata={
                "untrusted": True,
                "extraction": extraction,
                "size_bytes": len(data),
                "truncated": len(text) > self.max_text_chars,
                "redaction_count": int(redaction_events),
            },
        )

    @staticmethod
    def _decode_text(data: bytes) -> str:
        for encoding in ("utf-8-sig", "utf-8", "gb18030"):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue
        raise ValueError("text evidence encoding is not supported")

    def _extract_json(self, data: bytes, *, json_lines: bool) -> str:
        raw = self._decode_text(data)
        if json_lines:
            values = []
            for line_number, line in enumerate(raw.splitlines(), start=1):
                if not line.strip():
                    continue
                try:
                    values.append(json.loads(line))
                except json.JSONDecodeError as exc:
                    raise ValueError(f"invalid JSONL at line {line_number}") from exc
            value: Any = values
        else:
            try:
                value = json.loads(raw)
            except json.JSONDecodeError as exc:
                raise ValueError("invalid JSON evidence") from exc
        return json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True)

    def _extract_csv(self, data: bytes, *, delimiter: str) -> str:
        reader = csv.DictReader(io.StringIO(self._decode_text(data)), delimiter=delimiter)
        if not reader.fieldnames:
            raise ValueError("table evidence has no header")
        lines = []
        for index, row in enumerate(reader, start=1):
            if index > 5000:
                break
            cells = [f"{key}={value or ''}" for key, value in row.items() if key]
            lines.append(f"row {index}: " + " | ".join(cells))
        return "\n".join(lines)

    @staticmethod
    def _extract_pdf(data: bytes) -> str:
        try:
            from pypdf import PdfReader
        except ImportError as exc:  # pragma: no cover - optional install profile
            raise ValueError("PDF extraction dependency is unavailable") from exc
        try:
            reader = PdfReader(io.BytesIO(data), strict=False)
            if len(reader.pages) > 500:
                raise ValueError("PDF page limit exceeded")
            return "\n\n".join((page.extract_text() or "") for page in reader.pages)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("PDF evidence could not be parsed") from exc

    @classmethod
    def _extract_docx(cls, data: bytes) -> str:
        try:
            from docx import Document
        except ImportError as exc:  # pragma: no cover
            raise ValueError("Word extraction dependency is unavailable") from exc
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                entries = archive.infolist()
                if len(entries) > cls._MAX_DOCX_ENTRIES:
                    raise ValueError("DOCX archive entry limit exceeded")
                if sum(entry.file_size for entry in entries) > cls._MAX_DOCX_UNCOMPRESSED_BYTES:
                    raise ValueError("DOCX archive expansion limit exceeded")
                for entry in entries:
                    if entry.flag_bits & 0x1:
                        raise ValueError("encrypted DOCX evidence is not supported")
                    if (
                        entry.file_size >= 5 * 1024 * 1024
                        and entry.compress_size > 0
                        and entry.file_size / entry.compress_size > cls._MAX_DOCX_COMPRESSION_RATIO
                    ):
                        raise ValueError("DOCX archive expansion ratio exceeded")
            document = Document(io.BytesIO(data))
            parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(parts)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Word evidence could not be parsed") from exc

    def _extract_image(self, data: bytes) -> str:
        if self._image_extractor:
            return str(self._image_extractor(data))
        try:
            from PIL import Image
            import pytesseract

            image = Image.open(io.BytesIO(data))
            image.verify()
            image = Image.open(io.BytesIO(data))
            if image.width * image.height > 40_000_000:
                raise ValueError("image pixel limit exceeded")
            return str(pytesseract.image_to_string(image, lang="chi_sim+eng", timeout=10))
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("image OCR is unavailable or the image is invalid") from exc

    @staticmethod
    def _evidence_type(suffix: str) -> str:
        if suffix in {".log", ".json", ".jsonl", ".txt"}:
            return "audit_log"
        if suffix in {".csv", ".tsv"}:
            return "evaluation_dataset"
        if suffix in {".png", ".jpg", ".jpeg", ".webp", ".bmp"}:
            return "security_screenshot"
        return "security_document"

    @staticmethod
    def _default_media_type(suffix: str) -> str:
        return {
            ".json": "application/json",
            ".jsonl": "application/x-ndjson",
            ".csv": "text/csv",
            ".tsv": "text/tab-separated-values",
            ".pdf": "application/pdf",
            ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ".png": "image/png",
            ".jpg": "image/jpeg",
            ".jpeg": "image/jpeg",
        }.get(suffix, "text/plain")


__all__ = ["EvidenceProcessor", "ProcessedEvidence"]
